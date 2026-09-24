#!/usr/bin/env python3
"""
feladvanyok.xlsx  ->  assets/feladvanyok.js

Futtatas:      python scripts/xlsx_to_js.py
Teszt datum:   python scripts/xlsx_to_js.py --ma 2026-10-20
(A build.py magatol is lefuttatja, kulon nem kell.)

Mit csinal:
- beolvassa a munkafuzet "Feladványok" lapjat (a tobbi lapot nem nezi),
- ellenoriz mindent; ha barmi hibas, kiirja a feladvany azonositojat es a
  hibat, es LEALL, a JS-hez nem nyul,
- nem enged modositani mar elindult feladvanyt,
- a JS-be csak a mai es a kovetkezo ELORE_NAPOK nap feladvanyai kerulnek,
  mert a JS nyilvanos: a tavolabbi feladvanyok nem szivarognak ki.
A "csapda" oszlopot soha nem olvassa ki, igy az sehova nem kerulhet ki.
"""
import json, os, re, sys, warnings
from datetime import date, datetime, time, timedelta, timezone

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = os.path.join(ROOT, "feladvanyok.xlsx")
JS = os.path.join(ROOT, "assets", "feladvanyok.js")
SHEET = "Feladványok"
ELORE_NAPOK = 1   # ennyi jovobeli nap kerul a JS-be (a holnapi, hogy ejfelkor valthasson)
FIGYELMEZTETES_NAPOK = 7

KOTELEZO = ["feladvany_id", "kezdes", "szint", "kategoria", "szo1", "szo2", "szo3", "szo4"]
OSZLOPOK = KOTELEZO + ["megjegyzes"]  # a csapda szandekosan nincs itt


class BuildHiba(Exception):
    pass


# ---------------------------------------------------------------- datum
def budapest_now():
    """A pontos budapesti ido, a gep sajat idozonajatol fuggetlenul."""
    utc = datetime.now(timezone.utc)
    try:
        from zoneinfo import ZoneInfo
        return utc.astimezone(ZoneInfo("Europe/Budapest")).replace(tzinfo=None)
    except Exception:
        # nincs idozona-adatbazis (pl. Windows): EU nyari idoszamitas kezzel
        def last_sunday(month):
            d = date(utc.year, month, 31)
            return d - timedelta(days=(d.weekday() + 1) % 7)
        start = datetime.combine(last_sunday(3), time(1), tzinfo=timezone.utc)
        end = datetime.combine(last_sunday(10), time(1), tzinfo=timezone.utc)
        return (utc + timedelta(hours=2 if start <= utc < end else 1)).replace(tzinfo=None)


def today():
    """Budapesti mai datum, vagy a --ma kapcsoloval megadott teszt datum."""
    if "--ma" in sys.argv:
        return parse_date(sys.argv[sys.argv.index("--ma") + 1])
    return budapest_now().date()


def parse_date(s):
    if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", s):
        raise ValueError(s)
    y, m, d = map(int, s.split("-"))
    return date(y, m, d)


# ---------------------------------------------------------------- beolvasas
def load_js():
    """A mar kozzetett assets/feladvanyok.js tartalma."""
    if not os.path.exists(JS):
        return []
    raw = open(JS, encoding="utf-8").read()
    return json.loads(raw[raw.index("["):raw.rindex("]") + 1])


def cell_text(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def read_workbook(path=XLSX):
    """Beolvassa es ellenorzi a munkafuzetet. Hibanal BuildHiba-t dob, az osszes hibaval."""
    try:
        import openpyxl
    except ImportError:
        raise BuildHiba(["Hiányzik az openpyxl. Telepítsd: python -m pip install openpyxl"])
    warnings.filterwarnings("ignore", module="openpyxl")
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except PermissionError:
        raise BuildHiba(["A feladvanyok.xlsx most zárolva van (valószínűleg nyitva van az Excelben, "
                         "vagy a OneDrive épp szinkronizálja). Mentsd el, zárd be, és futtasd újra."])
    except Exception as e:
        raise BuildHiba(["Nem tudom megnyitni a feladvanyok.xlsx fájlt (%s)." % e])
    if SHEET not in wb.sheetnames:
        raise BuildHiba(['Nincs "%s" nevű lap a munkafüzetben.' % SHEET])
    rows = list(wb[SHEET].iter_rows(values_only=True))
    wb.close()
    if not rows:
        raise BuildHiba(['A "%s" lap üres.' % SHEET])

    header = [cell_text(h).lower() for h in rows[0]]
    missing = [c for c in OSZLOPOK if c not in header]
    if missing:
        raise BuildHiba(["Hiányzó oszlop(ok) az első sorban: %s" % ", ".join(missing)])
    col = {name: header.index(name) for name in OSZLOPOK}

    errors = []
    groups_by_id = {}
    order = []
    last_id = None
    for n, row in enumerate(rows[1:], start=2):
        get = lambda name: cell_text(row[col[name]]) if col[name] < len(row) else ""
        raw_kezdes = row[col["kezdes"]] if col["kezdes"] < len(row) else None
        if all(get(c) == "" for c in OSZLOPOK):
            continue  # teljesen ures sor
        pid = get("feladvany_id") or "(nincs azonosító)"
        where = "%s, Excel %d. sor" % (pid, n)
        for c in KOTELEZO:
            if get(c) == "":
                errors.append("%s: üres a(z) „%s” cella." % (where, c))
        if get("feladvany_id") and not re.fullmatch(r"[A-Za-z0-9_-]+", pid):
            errors.append("%s: az azonosítóban csak betű, szám, kötőjel és aláhúzás lehet." % where)
        level = int(get("szint")) if get("szint") in ("1", "2", "3", "4") else None
        if get("szint") and level is None:
            errors.append("%s: a szint csak 1, 2, 3 vagy 4 lehet (most: „%s”)." % (where, get("szint")))
        start = None
        if isinstance(raw_kezdes, datetime):
            start = raw_kezdes.date().isoformat()
        elif isinstance(raw_kezdes, date):
            start = raw_kezdes.isoformat()
        elif get("kezdes"):
            try:
                start = parse_date(get("kezdes")).isoformat()
            except ValueError:
                errors.append("%s: a kezdés nem érvényes dátum, ÉÉÉÉ-HH-NN formában kell (most: „%s”)." % (where, get("kezdes")))
        if not get("feladvany_id"):
            continue
        if pid in groups_by_id and last_id != pid:
            errors.append("%s: a(z) %s azonosító már szerepelt feljebb is. Két feladvány nem kaphatja ugyanazt az azonosítót." % (where, pid))
        last_id = pid
        if pid not in groups_by_id:
            groups_by_id[pid] = []
            order.append(pid)
        group = {"level": level, "title": get("kategoria"),
                 "words": [get("szo%d" % i) for i in range(1, 5)]}
        if get("megjegyzes"):
            group["note"] = get("megjegyzes")
        groups_by_id[pid].append((n, start, group))

    puzzles = []
    for pid in order:
        entries = groups_by_id[pid]
        if len(entries) != 4:
            errors.append("%s: %d sora van, pontosan 4 kell (szintenként egy)." % (pid, len(entries)))
        starts = {s for _, s, _ in entries if s}
        if len(starts) > 1:
            errors.append("%s: nem ugyanaz a kezdés dátuma mind a négy sorban (%s)." % (pid, ", ".join(sorted(starts))))
        levels = sorted(g["level"] for _, _, g in entries if g["level"])
        if len(entries) == 4 and levels != [1, 2, 3, 4] and len(levels) == 4:
            errors.append("%s: minden szintből (1–4) pontosan egy kell, most: %s." % (pid, ", ".join(map(str, levels))))
        words = [w.lower() for _, _, g in entries for w in g["words"] if w]
        for w in sorted({w for w in words if words.count(w) > 1}):
            errors.append("%s: a(z) „%s” szó többször is szerepel a feladványban." % (pid, w))
        puzzles.append({"id": pid, "start": min(starts) if starts else None,
                        "groups": sorted((g for _, _, g in entries), key=lambda g: g["level"] or 0)})
    if errors:
        raise BuildHiba(errors)
    return puzzles


# ---------------------------------------------------------------- osszefesules
def tartalom(p):
    """A feladvany osszehasonlithato tartalma (a szavak sorrendje nem szamit)."""
    return (p["start"], tuple((g["level"], g["title"], tuple(sorted(g["words"])), g.get("note", ""))
                              for g in sorted(p["groups"], key=lambda g: g["level"])))


def merge(workbook, published, ma):
    """A munkafuzet + a mar kozzetett feladvanyok. Elindult feladvanyt nem enged modositani."""
    errors = []
    wb_by_id = {str(p["id"]): p for p in workbook}
    result = {}
    kept_old, dropped = [], []
    for p in published:
        key = str(p["id"])
        live = parse_date(p["start"]) <= ma
        if not live:
            if key not in wb_by_id:
                dropped.append(key)
            continue  # meg nem indult el, a munkafuzet donti el, mi lesz vele
        if key in wb_by_id and tartalom(wb_by_id[key]) != tartalom(p):
            errors.append("%s: ez a feladvány már élesben ment (%s), de a munkafüzetben megváltozott. "
                          "Élő feladványt nem szabad módosítani: állítsd vissza az eredeti szavakat, "
                          "kategóriákat és dátumot, vagy adj új azonosítót." % (key, p["start"]))
        if key not in wb_by_id:
            kept_old.append(key)
        result[key] = p  # mindig a kozzetett valtozat marad
    for p in workbook:
        key = str(p["id"])
        if key not in result:
            result[key] = p
    by_date = {}
    for p in result.values():
        by_date.setdefault(p["start"], []).append(str(p["id"]))
    for day, ids in sorted(by_date.items()):
        if len(ids) > 1:
            errors.append("%s: ezen a napon több feladvány is indulna (%s). Naponta csak egy lehet."
                          % (day, ", ".join(sorted(ids))))
    if errors:
        raise BuildHiba(errors)
    return sorted(result.values(), key=lambda p: p["start"]), kept_old, dropped


def remaining_days(puzzles, ma):
    """Hany egymast koveto napra van feladvany mattol kezdve, es mikor jon a kovetkezo utana."""
    days = {p["start"] for p in puzzles}
    n, d = 0, ma
    while d.isoformat() in days:
        n += 1
        d += timedelta(days=1)
    later = sorted(s for s in days if s > d.isoformat())
    return n, (d - timedelta(days=1)), (later[0] if later else None)


def report(puzzles, ma):
    """Kiirja, meddig van feladvany. True, ha fogyoban vannak (a tervezett szunet nem szamit)."""
    n, last, nxt = remaining_days(puzzles, ma)
    if n == 0:
        line = "ma (%s) nincs új feladvány, a legutóbbi jelenik meg." % ma
    else:
        line = "Feladvány még %d napra van, %s-ig." % (n, last)
    if nxt:
        line += " Utána a következő: %s." % nxt
    elif n:
        line += " Utána nincs több feladvány a munkafüzetben."
    alarm = n < FIGYELMEZTETES_NAPOK and not (n == 0 and nxt)
    print(("FIGYELEM: " if alarm else "") + line[0].upper() + line[1:] + (" Ideje újakat írni." if alarm else ""))
    return alarm


# ---------------------------------------------------------------- kiiras
HEADER = """/* =====================================================================
   FELADVÁNYOK – ezt a fájlt a build.py generálja a feladvanyok.xlsx-ből.
   NE szerkeszd kézzel: a munkafüzetet írd át, és futtasd a build.py-t.
   - start: az a nap, amikor a feladvány élesedik (ÉÉÉÉ-HH-NN, budapesti idő)
   - level: 1 = legkönnyebb ... 4 = legtrükkösebb
   - note: rövid kiegészítés, ami a megfejtés után jelenik meg
   ===================================================================== */
"""


def to_js(puzzles):
    def dump(v):
        return json.dumps(v, ensure_ascii=False)
    blocks = []
    for p in puzzles:
        groups = []
        for g in sorted(p["groups"], key=lambda g: g["level"]):
            line = '      { "level": %d, "title": %s, "words": %s' % (g["level"], dump(g["title"]), dump(g["words"]))
            if g.get("note"):
                line += ', "note": %s' % dump(g["note"])
            groups.append(line + " }")
        blocks.append('  {\n    "id": %s,\n    "start": %s,\n    "groups": [\n%s\n    ]\n  }'
                      % (dump(p["id"]), dump(p["start"]), ",\n".join(groups)))
    return HEADER + "window.FELADVANYOK = [\n" + ",\n".join(blocks) + "\n];\n"


def run(ma=None):
    """A build.py innen hivja. Visszaadja az osszes (a jovobelieket is tartalmazo) feladvanyt."""
    ma = ma or today()
    published = load_js()
    if not os.path.exists(XLSX):
        print("Nincs feladvanyok.xlsx, a meglévő assets/feladvanyok.js marad.")
        return published, False
    workbook = read_workbook()
    puzzles, kept_old, dropped = merge(workbook, published, ma)
    limit = (ma + timedelta(days=ELORE_NAPOK)).isoformat()
    public = [p for p in puzzles if p["start"] <= limit]
    content = to_js(public)
    old = open(JS, encoding="utf-8").read() if os.path.exists(JS) else ""
    if content != old:
        open(JS, "w", encoding="utf-8", newline="\n").write(content)
        print("   assets/feladvanyok.js frissítve")
    print("Munkafüzet: %d feladvány, a JS-be került: %d (%s-ig)" % (len(workbook), len(public), limit))
    if kept_old:
        print("   korábbi, csak a JS-ben meglévő feladvány(ok) megtartva: %s" % ", ".join(kept_old))
    if dropped:
        print("   még el nem indult, a munkafüzetben nem szereplő feladvány(ok) kivéve: %s" % ", ".join(dropped))
    return puzzles, True


def hibak_kiirasa(errors):
    print()
    print("A BUILD LEÁLLT, semmit nem írtam át. Javítsd ezeket a munkafüzetben:")
    for e in errors:
        print("  - " + e)
    print()


if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    try:
        ma = today()
        puzzles, from_xlsx = run(ma)
        alarm = report(puzzles, ma)
        # --keszlet: a GitHub Action ezzel jelez (hibás futás = e-mail), ha fogyóban vannak a feladványok
        if "--keszlet" in sys.argv and from_xlsx and alarm:
            sys.exit(2)
    except BuildHiba as e:
        hibak_kiirasa(e.args[0])
        sys.exit(1)
